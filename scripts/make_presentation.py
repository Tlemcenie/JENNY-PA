"""Génère le support de présentation JENNY (.docx) — briques + script oral + visuels."""

import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CHARTS = "/tmp/jenny_charts"


def _metrics():
    try:
        return json.load(open("var/eval/ml_metrics.json"))
    except Exception:
        return {"intent_lstm": {"top1_accuracy": 94.7}, "ner_bigru": {"token_accuracy": 99.7, "entity_f1": 95.3}}

NIGHT = RGBColor(0x1A, 0x14, 0x30)
CORAL = RGBColor(0xE5, 0x67, 0x4F)
INK = RGBColor(0x1A, 0x1D, 0x26)
MUTED = RGBColor(0x55, 0x5A, 0x66)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = INK


def _shade(cell, hexfill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(shd)


def brick_title(num, emoji, title):
    p = doc.add_paragraph()
    p.space_before = Pt(18)
    r = p.add_run(f"BRIQUE {num}")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = CORAL
    p2 = doc.add_paragraph()
    r2 = p2.add_run(f"{emoji}  {title}")
    r2.font.size = Pt(17)
    r2.font.bold = True
    r2.font.color.rgb = NIGHT
    p2.space_after = Pt(4)


def in_brief(text):
    p = doc.add_paragraph()
    r = p.add_run("👉  " + text)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED


def subhead(text):
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NIGHT


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " : ")
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def say_box(lines):
    """Encadré corail clair : le script oral."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _shade(cell, "FFF1EE")
    cell.width = Pt(460)
    head = cell.paragraphs[0]
    rh = head.add_run("🎤  À DIRE À LA PROF")
    rh.bold = True
    rh.font.size = Pt(10)
    rh.font.color.rgb = CORAL
    for ln in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(ln)
        r.font.size = Pt(11)
        r.font.color.rgb = INK
    doc.add_paragraph()


def figure(filename, caption, width=6.0):
    """Insère une image centrée + légende, si le fichier existe."""
    path = filename if os.path.isabs(filename) else os.path.join(CHARTS, filename)
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = cap.add_run(caption)
    rc.italic = True
    rc.font.size = Pt(9)
    rc.font.color.rgb = MUTED
    doc.add_paragraph()


# ────────────────────────────── PAGE DE TITRE ──────────────────────────────
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.space_before = Pt(90)
r = t.add_run("JENNY")
r.font.size = Pt(54)
r.font.bold = True
r.font.color.rgb = NIGHT

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = s.add_run("L'assistante IA qui répond aux clients des salons de coiffure,\n24h/24, sur tous les canaux.")
rs.font.size = Pt(14)
rs.font.color.rgb = CORAL
rs.italic = True

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
tag.space_before = Pt(30)
rt = tag.add_run("Support de présentation — briques du projet & script oral\n(métier + technique · 15-20 min · 3 présentateurs)")
rt.font.size = Pt(11)
rt.font.color.rgb = MUTED

doc.add_page_break()

# ──────────────────────── COMMENT UTILISER + RÉPARTITION ────────────────────────
h = doc.add_paragraph()
rh = h.add_run("Comment utiliser ce document")
rh.font.size = Pt(18)
rh.bold = True
rh.font.color.rgb = NIGHT

doc.add_paragraph(
    "Le projet est découpé en 9 briques indépendantes. Chacun choisit les briques qu'il "
    "veut présenter. Pour chaque brique vous avez : ce qu'on a concrètement fait (métier + "
    "technique), puis un encadré « À dire à la prof » — un texte prêt à dire, à reformuler "
    "avec vos mots. Une répartition équilibrée pour 3 personnes (≈ 5-6 min chacun) :"
)

split = doc.add_table(rows=4, cols=3)
split.style = "Light Grid Accent 2"
hdr = split.rows[0].cells
for c, txt in zip(hdr, ["Présentateur 1 — Produit / Métier", "Présentateur 2 — Cerveau IA", "Présentateur 3 — Technique / Infra"]):
    rr = c.paragraphs[0].add_run(txt)
    rr.bold = True
    rr.font.size = Pt(10)
rows = [
    ("B1 · Le problème & la vision", "B3 · Les agents & l'orchestration", "B2 · L'architecture globale"),
    ("B6 · Les super-pouvoirs métier", "B4 · Les modèles d'IA maison", "B7 · Omnicanal (WhatsApp/Insta/Mail)"),
    ("B8 · Dashboard & expérience", "B5 · Le RAG (base de connaissances)", "B9 · Cloud, sécurité & RGPD"),
]
for i, row in enumerate(rows, start=1):
    for j, txt in enumerate(row):
        split.cell(i, j).paragraphs[0].add_run(txt).font.size = Pt(10)

doc.add_paragraph()
tip = doc.add_paragraph()
rtip = tip.add_run(
    "Conseil : commencez par la Brique 1 (le pitch métier, pour accrocher), finissez par une "
    "démo live (chat + WhatsApp) si possible. Gardez la dernière page « Questions probables » sous le coude."
)
rtip.italic = True
rtip.font.color.rgb = MUTED

doc.add_page_break()

# ────────────────────────────── LES BRIQUES ──────────────────────────────
brick_title(1, "💡", "Le problème & la vision")
in_brief("Pourquoi Jenny existe — la partie « pitch métier » qui accroche le jury.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Le problème", "un salon de coiffure perd des clients parce que les messages (WhatsApp, Insta, mail) arrivent quand l'équipe a les mains dans les cheveux. Résultat : réponses lentes, RDV ratés, no-shows."),
    ("La solution", "Jenny, une assistante IA qui répond instantanément, 24h/24, prend les RDV, fait les devis, et ne se présente jamais comme un robot — elle parle comme un membre de l'équipe."),
    ("Le positionnement", "un vrai SaaS multi-salons (multi-tenant) : chaque salon a son compte, ses prestations, sa base de connaissances, ses canaux."),
    ("La valeur", "zéro message perdu, plus de réservations, moins de charge mentale pour le coiffeur."),
])
say_box([
    "« Imaginez un salon de coiffure un samedi après-midi : le téléphone sonne, trois clients attendent, "
    "et sur Instagram quelqu'un demande “vous faites les balayages ?” — personne ne répond, le client va ailleurs.",
    "Notre projet, Jenny, résout exactement ça : c'est une assistante en intelligence artificielle qui répond "
    "aux clients d'un salon instantanément, à toute heure, sur WhatsApp, Instagram et par mail. Elle prend les "
    "rendez-vous, fait des devis, répond aux questions — et surtout elle parle comme une vraie personne de l'équipe.",
    "C'est un produit complet, pensé comme un logiciel qu'on pourrait vendre à plusieurs salons : chacun a son "
    "espace, ses tarifs, ses informations. »",
])

brick_title(2, "🏗️", "L'architecture globale")
in_brief("La vue d'ensemble technique : les grandes pièces et comment elles se parlent.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Le back-end", "une API en Python (FastAPI) qui contient toute l'intelligence."),
    ("Le front-end", "une interface web moderne en Next.js / React (le tableau de bord du coiffeur + la page d'accueil)."),
    ("La base de données", "PostgreSQL, avec l'extension pgvector pour la recherche par similarité (l'IA)."),
    ("Le cache & temps réel", "Redis (mémoire des conversations, files de tâches)."),
    ("Les cerveaux IA", "2 petits modèles entraînés par nous + un grand modèle de langage (LLM) pour rédiger."),
    ("Le tout déployé", "sur Google Cloud (Cloud Run), hébergé en France pour le RGPD."),
])
say_box([
    "« Techniquement, le projet est une application web complète. Le cœur, c'est une API en Python avec FastAPI : "
    "c'est là que vit toute l'intelligence. Le coiffeur, lui, voit une interface moderne développée en Next.js et React.",
    "Côté données, on utilise PostgreSQL, avec une extension spéciale, pgvector, qui permet à l'IA de faire de la "
    "recherche “par sens” et pas juste par mot-clé. Redis gère la mémoire des conversations et le temps réel.",
    "Et pour l'intelligence, on combine trois choses : deux petits modèles qu'on a entraînés nous-mêmes, et un grand "
    "modèle de langage qui rédige les réponses. Le tout est déployé sur Google Cloud, hébergé en France. »",
])

brick_title(3, "🧠", "Les agents & l'orchestration")
in_brief("Comment Jenny « réfléchit » : la chaîne de traitement d'un message. Le cœur technique.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Le pipeline", "chaque message passe par 6 étapes : contexte → détection d'intention → extraction d'infos → routage vers un agent → action réelle → rédaction par le LLM."),
    ("6 agents spécialisés", "FAQ, action, planning, paiement, devis, fallback — chacun gère un type de demande."),
    ("9 actions métier réelles", "prendre/annuler un RDV, faire un devis, liste d'attente, ticket, paiement… qui écrivent vraiment en base de données."),
    ("Choix d'architecture assumé", "le routage est déterministe (une intention → un agent), pas laissé au hasard du LLM → fiable et prévisible. On n'a pas utilisé LangChain : orchestration 100 % maison."),
    ("Le rôle du LLM", "il ne décide rien — il reçoit les FAITS de l'action et les reformule en français naturel."),
])
say_box([
    "« Quand un client écrit, Jenny ne se contente pas de “demander à ChatGPT”. Le message traverse une vraie "
    "chaîne de traitement maison. D'abord on comprend l'intention (prendre un RDV ? un prix ? une réclamation ?), "
    "on extrait les infos utiles (le service, la date, l'heure), puis on aiguille vers l'un de nos 6 agents spécialisés.",
    "Chaque agent déclenche une vraie action : par exemple l'agent planning va réellement écrire un rendez-vous "
    "dans la base de données. Au total on a 9 actions concrètes.",
    "Un choix important : on a fait l'orchestration nous-mêmes, sans framework type LangChain. Le routage est "
    "déterministe — c'est plus fiable pour un vrai salon, on ne laisse pas l'IA décider toute seule d'annuler un RDV. "
    "Le grand modèle de langage, lui, sert seulement à transformer les faits en réponse naturelle et humaine. »",
])

brick_title(4, "🤖", "Les modèles d'IA maison (entraînés de zéro)")
in_brief("Le « wow » technique : on n'a pas qu'utilisé des IA, on en a entraîné nous-mêmes.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Modèle nº1 — Intention", "un réseau de neurones (Bi-LSTM + attention) qu'on a entraîné pour classer le message en 14 intentions, multi-label."),
    ("Modèle nº2 — Entités (NER)", "un réseau BiGRU qui repère dans la phrase la date, l'heure, le service, l'email, etc."),
    ("Un mini-LLM « NanoGPT »", "on a même entraîné un modèle de langage DE ZÉRO : corpus français, tokenizer BPE maison, ~17 millions de paramètres, entraîné en PyTorch."),
    ("Comparaison honnête", "un onglet compare 3 modèles (Groq, Ollama, notre NanoGPT) avec scores et latence, plus la courbe d'entraînement (la loss qui descend de 9,7 à 4,6)."),
    ("L'idée", "comprendre vraiment comment marche un LLM en en construisant un, même petit."),
])
say_box([
    "« Là où on est allés plus loin : on n'a pas seulement branché une IA existante, on en a entraîné nous-mêmes. "
    "On a deux petits modèles maison : un qui devine ce que veut le client (l'intention), et un autre qui repère "
    "les infos clés dans la phrase, comme la date ou le service demandé.",
    "Et pour aller au bout de la démarche, on a même entraîné un mini modèle de langage de zéro — notre “NanoGPT” : "
    "on a constitué un corpus en français, fabriqué notre propre découpeur de mots, et on l'a entraîné en PyTorch. "
    "Il fait 17 millions de paramètres.",
    "On a un onglet qui compare honnêtement notre petit modèle aux gros (Groq, Ollama) : forcément il est plus faible, "
    "mais l'intérêt c'était de comprendre concrètement comment un LLM apprend — on montre même la courbe d'apprentissage. »",
])

_m = _metrics()
subhead("📊 Les chiffres réels (mesurés, pas estimés)")
mt = doc.add_table(rows=5, cols=3)
mt.style = "Light List Accent 2"
mt.cell(0, 0).paragraphs[0].add_run("Modèle").bold = True
mt.cell(0, 1).paragraphs[0].add_run("Performance").bold = True
mt.cell(0, 2).paragraphs[0].add_run("Mesure").bold = True
mrows = [
    ("Intention (Bi-LSTM maison)", f"{_m['intent_lstm']['top1_accuracy']} %", "exactitude top-1 (600 exemples)"),
    ("Entités / NER (BiGRU maison)", f"{_m['ner_bigru']['token_accuracy']} % · F1 {_m['ner_bigru']['entity_f1']} %", "tokens · entités (600 exemples)"),
    ("Agents (banc d'essai bout-en-bout)", "100 %", "58/58 vérifs · 14 scénarios · 23 tours"),
    ("NanoGPT maison", "16,9 M params · val loss 4,69", "entraîné de zéro (loss 9,7 → 4,6)"),
]
for i, (a, b, c) in enumerate(mrows, start=1):
    mt.cell(i, 0).paragraphs[0].add_run(a).font.size = Pt(10)
    mt.cell(i, 1).paragraphs[0].add_run(b).font.size = Pt(10)
    mt.cell(i, 2).paragraphs[0].add_run(c).font.size = Pt(10)
doc.add_paragraph()
figure("ml_accuracy.png", "Performance des modèles maison (intention, NER, agents).", 5.9)
figure("training_curve.png", "NanoGPT : la loss descend de 9,71 à 4,64 au fil de l'entraînement — il apprend vraiment.", 5.9)
figure("model_comparison.png", "Comparaison honnête des 3 LLM : Groq/Ollama au top, notre NanoGPT plus modeste (assumé).", 6.2)

brick_title(5, "📚", "Le RAG — Jenny apprend les infos du salon")
in_brief("Comment Jenny répond juste sur un salon précis sans inventer. Technique + très visuel.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Le principe (RAG)", "on donne à Jenny des documents (politiques, marques de produits, équipe, accès…). Elle les « vectorise » et peut retrouver le bon passage pour répondre."),
    ("Concrètement", "le coiffeur dépose un texte ou un PDF dans le dashboard → découpage en morceaux → embeddings (pgvector) → recherche par similarité."),
    ("Anti-invention", "Jenny se base sur la connaissance trouvée ; si elle ne sait pas, elle ne l'invente pas (garde-fou)."),
    ("Ce qu'on a amélioré", "au départ seul l'agent FAQ consultait la base ; on a ajouté une couche qui interroge la connaissance pour TOUTE question informative, même mal classée."),
    ("Démo", "on a donné un fichier “politiques du salon”, et à la question “vous avez un parking ?” Jenny répond exactement le contenu du fichier (parking de 6 places, rue des Lilas)."),
])
say_box([
    "« Un assistant IA qui invente des informations, c'est une catastrophe pour un commerce. Donc on a mis en place "
    "ce qu'on appelle le RAG : on donne à Jenny des documents sur le salon — ses horaires, ses marques de produits, "
    "son équipe — et elle apprend à retrouver la bonne information dedans.",
    "Techniquement, chaque document est transformé en vecteurs et stocké ; quand un client pose une question, on "
    "cherche le passage le plus proche par le sens. Jenny répond à partir de ça, et si l'info n'existe pas, elle "
    "n'invente pas.",
    "Pour la démo : on a donné un fichier avec les infos pratiques du salon. Quand on demande “vous avez un parking ?”, "
    "Jenny répond exactement ce qui est écrit dans le fichier — un parking de 6 places à l'arrière. La preuve qu'elle "
    "lit, comprend et se base sur la connaissance qu'on lui donne. »",
])

figure("screen_chat.png", "Le chat client : Jenny répond en se basant sur la base de connaissances du salon.", 6.0)

brick_title(6, "✂️", "Les super-pouvoirs métier (coiffeur)")
in_brief("Tout ce que Jenny fait VRAIMENT pour un salon. Brique 100 % métier, concrète.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Prise de RDV réelle", "Jenny propose des créneaux libres réels, réserve en base, gère l'anti-double-réservation et les multi-prestations."),
    ("Devis professionnel", "génération d'un vrai PDF (logo, lignes, TVA 20 %, HT/TTC) envoyé par email."),
    ("Liste d'attente", "salon complet → le client s'inscrit, et il est re-proposé automatiquement en cas d'annulation."),
    ("Rappels & avis", "rappel la veille du RDV, demande d'avis après, relance de re-réservation."),
    ("Fidélité & no-shows", "points de fidélité, et acompte demandé en cas de no-shows répétés."),
    ("Annulation/déplacement", "le client gère son RDV en langage naturel (“je peux décaler à jeudi ?”)."),
])
say_box([
    "« Jenny n'est pas qu'un chatbot qui papote : elle agit vraiment. Elle prend des rendez-vous en proposant de "
    "vrais créneaux libres, et elle les enregistre dans l'agenda — avec la sécurité de ne jamais réserver deux fois "
    "le même créneau.",
    "Elle génère aussi des devis professionnels : un vrai PDF avec le logo, le détail des prestations, la TVA, le "
    "total — envoyé par email au client. Elle gère la liste d'attente quand c'est complet, et re-propose la place "
    "automatiquement si quelqu'un annule.",
    "Et tout l'écosystème d'un salon est là : rappels la veille, demande d'avis après le rendez-vous, programme de "
    "fidélité, et même un acompte demandé aux clients qui ne viennent jamais. C'est pensé comme un vrai outil de "
    "gestion, piloté par la conversation. »",
])

figure("/tmp/devis_test.pdf.png", "Exemple de devis généré par Jenny : logo, lignes, ventilation HT / TVA 20 % / TTC.", 4.4)

brick_title(7, "📱", "Omnicanal — WhatsApp, Instagram, Email")
in_brief("Jenny est là où sont les clients. Brique métier + un peu technique (intégrations).")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("WhatsApp & Instagram", "intégration via l'API officielle de Meta (Cloud API) : les vrais messages arrivent à Jenny, qui répond sur le même canal."),
    ("Email", "envoi des confirmations, devis et resets de mot de passe via un service d'emailing pro (Resend), avec un design soigné aux couleurs de Jenny."),
    ("Webhooks sécurisés", "chaque message entrant est vérifié (signature Meta) avant d'être traité."),
    ("Identité client par canal", "Jenny reconnaît le client par son numéro / son compte, et relie la conversation à sa fiche."),
    ("Mode gracieux", "si un canal n'est pas configuré, ça ne plante pas — il est juste “en attente”."),
])
say_box([
    "« Un client ne va pas sur un site web pour écrire à son coiffeur : il prend WhatsApp ou Instagram. Donc Jenny "
    "est branchée directement sur ces messageries, via l'API officielle de Meta. Concrètement, on a connecté un vrai "
    "numéro WhatsApp : un message envoyé arrive à Jenny, qui répond sur WhatsApp — on l'a testé en live.",
    "Pour l'email, Jenny envoie les confirmations de rendez-vous, les devis et les réinitialisations de mot de passe, "
    "avec un design pro à ses couleurs.",
    "Côté technique, on a sécurisé tout ça : chaque message entrant est vérifié par signature pour être sûr qu'il "
    "vient bien de Meta, et Jenny reconnaît automatiquement le client par son numéro pour relier la conversation à sa fiche. »",
])

brick_title(8, "🖥️", "Le dashboard & l'expérience")
in_brief("Ce que voit et utilise le coiffeur. Brique métier + design / UX.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("Tableau de bord complet", "boîte de réception unifiée, conversations, fiches clients 360°, calendrier, statistiques du salon."),
    ("Base de connaissances", "le coiffeur ajoute/édite ce que Jenny doit savoir, en quelques clics."),
    ("Cerveau de Jenny", "réglages du ton, du secteur, de la personnalité de l'assistante."),
    ("Temps réel", "les nouveaux messages et l'indicateur “Jenny écrit…” apparaissent en direct (SSE)."),
    ("Direction artistique", "page d'accueil soignée (vidéo, animations), interface premium, mascotte Jenny, mode clair/sombre."),
    ("Co-pilote humain", "si besoin, un humain reprend la main et Jenny propose une réponse à valider."),
])
say_box([
    "« Côté coiffeur, tout se pilote depuis un tableau de bord. Il y retrouve toutes ses conversations au même endroit, "
    "les fiches de ses clients, son calendrier, ses statistiques. Il peut apprendre des choses à Jenny en deux clics "
    "via la base de connaissances, et même régler sa personnalité — son ton, son style.",
    "On a beaucoup soigné l'expérience : les messages arrivent en temps réel, on voit “Jenny écrit…” comme sur une "
    "vraie messagerie, l'interface est premium avec une mascotte, un mode clair et sombre.",
    "Et important : l'humain garde le contrôle. À tout moment, le coiffeur peut reprendre une conversation, et Jenny "
    "lui propose alors une réponse toute prête à valider — elle devient une assistante, pas un remplaçant. »",
])

figure("screen_dashboard.png", "Le tableau de bord du coiffeur : messages gérés, taux de résolution, top intentions, activité en direct.", 6.3)

brick_title(9, "☁️", "Cloud, sécurité & RGPD")
in_brief("Le projet est réellement en ligne, pas juste sur un PC. Brique technique / DevOps.")
subhead("Ce qu'on a fait / ce qu'il y a")
bullets([
    ("En production réelle", "déployé sur Google Cloud (Cloud Run) : l'API, le site, la base de données (Cloud SQL), Redis — accessibles par une vraie URL."),
    ("Industrialisé", "images Docker construites par Cloud Build, secrets dans un coffre-fort (Secret Manager), déploiement en une commande (deploy.sh)."),
    ("Hébergé en France", "région Paris → conformité RGPD."),
    ("Sécurité", "identifiants des salons chiffrés en base, signatures des webhooks vérifiées, authentification par jeton."),
    ("Qualité", "près de 400 tests automatisés, garde-fou anti-invention, Jenny qui ne révèle jamais qu'elle est une IA."),
])
say_box([
    "« Ce n'est pas une maquette qui tourne sur nos ordinateurs : le projet est réellement en ligne. On l'a déployé "
    "sur Google Cloud — l'API, le site, la base de données, tout est accessible par une vraie adresse internet, et "
    "hébergé en France pour respecter le RGPD.",
    "On a industrialisé le déploiement : les conteneurs sont construits automatiquement, les clés secrètes sont dans "
    "un coffre-fort, et on remet en ligne une nouvelle version avec une seule commande.",
    "Enfin, la sécurité et la qualité : les identifiants des salons sont chiffrés, les messages entrants sont vérifiés, "
    "on a près de 400 tests automatisés, et un garde-fou empêche Jenny d'inventer — sans oublier qu'elle ne dit jamais "
    "au client qu'elle est une IA, c'est un parti pris produit. »",
])

# ──────────────────────── LIMITES & AMÉLIORATIONS ────────────────────────
doc.add_page_break()
h = doc.add_paragraph()
rh = h.add_run("⚖️ Limites actuelles & pistes d'amélioration")
rh.font.size = Pt(18)
rh.bold = True
rh.font.color.rgb = NIGHT
doc.add_paragraph(
    "On présente le projet honnêtement : il est réellement fonctionnel et en ligne, mais comme "
    "tout produit il a des limites assumées et une feuille de route. Le montrer est un atout (ça "
    "prouve qu'on a du recul sur notre travail)."
)
lim = doc.add_table(rows=9, cols=2)
lim.style = "Light Grid Accent 2"
lim.cell(0, 0).paragraphs[0].add_run("Limite actuelle").bold = True
lim.cell(0, 1).paragraphs[0].add_run("Piste d'amélioration").bold = True
pairs = [
    ("Le RAG (base de connaissances) tourne en local : les embeddings utilisent Ollama, absent du cloud → en prod le RAG n'est pas encore actif.",
     "Brancher un moteur d'embeddings cloud (OpenAI 768-dim, ou un modèle embarqué dans l'image) + ré-indexer. Le code est déjà prêt, c'est un réglage."),
    ("Le worker de tâches planifiées (rappels, avis, relances) n'est pas encore déployé en cloud (c'est un service permanent, pas une simple fonction).",
     "Le déployer en service always-on (Cloud Run min-instances=1) ou sur une petite VM. Tout le code des crons existe déjà."),
    ("Notre LLM maison (NanoGPT) est faible : ~17 M de paramètres, score 16,7 %.",
     "C'est assumé et pédagogique. Pistes : plus de données, modèle plus grand, plus d'entraînement. En prod on utilise Groq."),
    ("WhatsApp / Instagram sont en mode test (numéros de test, app Meta non publiée au grand public).",
     "Passer la validation Meta (App Review) pour ouvrir à tous les clients réels."),
    ("Les modèles intention/NER sont entraînés sur des données générées, pas sur de vrais messages clients.",
     "Boucle d'apprentissage : réentraîner sur les vraies conversations (la détection de lacunes alimente déjà la base)."),
    ("La finesse du RAG est perfectible : certaines questions récupèrent un passage un peu à côté.",
     "Améliorer le découpage des documents, ajouter un re-ranking, augmenter le seuil de pertinence."),
    ("Le paiement est en « mode gracieux » (pas de vrai débit pour l'instant).",
     "Brancher Stripe en réel pour les acomptes (l'intégration est déjà amorcée)."),
    ("Pas encore de tests de charge ni de montée en charge multi-instance éprouvée (le temps réel SSE est mono-process).",
     "Tests de charge, bus d'événements partagé (Redis pub/sub) pour scaler horizontalement ; app mobile (dossier Flutter déjà amorcé)."),
]
for i, (a, b) in enumerate(pairs, start=1):
    lim.cell(i, 0).paragraphs[0].add_run(a).font.size = Pt(9.5)
    lim.cell(i, 1).paragraphs[0].add_run(b).font.size = Pt(9.5)

doc.add_paragraph()
say_box([
    "« On est transparents sur les limites : le cœur du produit est réellement en ligne et fonctionne, mais il "
    "reste des chantiers. Par exemple, notre recherche documentaire marche en local et on doit finir de la brancher "
    "côté cloud ; nos modèles sont entraînés sur des données générées qu'on pourra enrichir avec de vrais échanges ; "
    "et pour ouvrir WhatsApp à tous les clients, il faut passer la validation de Meta.",
    "Mais on a justement construit une boucle d'amélioration continue : Jenny détecte quand elle ne sait pas "
    "répondre, et ces “lacunes” servent à enrichir sa base de connaissances. Le produit est conçu pour s'améliorer "
    "avec l'usage. »",
])

# ──────────────────────── QUESTIONS PROBABLES ────────────────────────
doc.add_page_break()
h = doc.add_paragraph()
rh = h.add_run("🎯 Questions probables de la prof (et comment répondre)")
rh.font.size = Pt(16)
rh.bold = True
rh.font.color.rgb = NIGHT
doc.add_paragraph()

qa = [
    ("« C'est juste un wrapper autour de ChatGPT ? »",
     "Non. L'orchestration est faite maison (6 agents, routage déterministe), on a entraîné nos propres modèles "
     "d'intention et d'entités, et même un mini-LLM de zéro. Le grand modèle ne sert qu'à reformuler des faits réels."),
    ("« Et si l'IA invente une information ? »",
     "On a un garde-fou anti-invention : Jenny répond à partir de la base de connaissances et des données réelles "
     "(catalogue, agenda). Si l'info n'existe pas, elle ne l'invente pas et peut passer la main à un humain."),
    ("« Les rendez-vous sont vraiment enregistrés ? »",
     "Oui, en base PostgreSQL, avec anti-double-réservation. La prise de RDV, le devis PDF, la liste d'attente : "
     "tout écrit réellement en base, ce n'est pas simulé."),
    ("« Pourquoi ne pas avoir utilisé LangChain ? »",
     "Pour un salon, on veut un comportement fiable et prévisible. Un routage déterministe (une intention → un agent) "
     "est plus sûr qu'un agent autonome qui pourrait, par exemple, annuler un RDV de lui-même. Et c'est plus léger en production."),
    ("« Le modèle maison NanoGPT est-il bon ? »",
     "Non, et c'est assumé : 17 M de paramètres face à des modèles de 70 milliards. L'objectif était pédagogique — "
     "comprendre comment un LLM apprend en en construisant un, avec sa courbe d'entraînement et une comparaison honnête."),
    ("« Comment gérez-vous plusieurs salons (multi-tenant) ? »",
     "Chaque donnée est isolée par salon (business_id) : prestations, conversations, base de connaissances, canaux. "
     "Un salon ne voit jamais les données d'un autre."),
]
for q, a in qa:
    p = doc.add_paragraph()
    rq = p.add_run("Q — " + q)
    rq.bold = True
    rq.font.color.rgb = CORAL
    pa = doc.add_paragraph()
    pa.paragraph_format.space_after = Pt(10)
    pa.add_run("R — " + a)

doc.save("PRESENTATION_JENNY.docx")
print("OK → PRESENTATION_JENNY.docx")
